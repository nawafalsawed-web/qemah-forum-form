#!/usr/bin/env python3
"""تقرير زوار جناح قمة للمزادات: python3 report.py  ->  تقرير-زوار-الجناح.xlsx + ملخص على الشاشة"""
import json, urllib.request, urllib.parse, collections
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

PROJECT = 'ahdaf-influencers'; COL = 'qemah_visitors'
CID = '563584335869-fgrhgmd47bqnekij5i8b5pr03ho849e6.apps.googleusercontent.com'; CSEC = 'j9iVZfS8kkCEFUPaAeJV0sAi'
FIELDS = [('name','الاسم'),('phone','الجوال'),('title','المسمى الوظيفي'),('org','الجهة'),('category','الفئة'),
          ('interests','الاهتمامات'),('note','الاستفسار'),('consent','موافقة المتابعة'),('usher','المنظّم'),('day','اليوم'),('ts','وقت التسجيل')]

def token():
    import os
    rt = json.load(open(os.path.expanduser('~/.config/configstore/firebase-tools.json')))['tokens']['refresh_token']
    d = urllib.parse.urlencode({'client_id':CID,'client_secret':CSEC,'refresh_token':rt,'grant_type':'refresh_token'}).encode()
    return json.load(urllib.request.urlopen('https://oauth2.googleapis.com/token', d))['access_token']

def val(v):
    k = next(iter(v)); x = v[k]
    if k == 'arrayValue': return '، '.join(val(i) for i in x.get('values', []))
    if k == 'booleanValue': return 'نعم' if x else 'لا'
    return str(x)

def fetch(at):
    docs, tok = [], ''
    while True:
        url = f'https://firestore.googleapis.com/v1/projects/{PROJECT}/databases/(default)/documents/{COL}?pageSize=300&pageToken={tok}'
        r = json.load(urllib.request.urlopen(urllib.request.Request(url, headers={'Authorization':'Bearer '+at})))
        docs += [{k: val(v) for k, v in d.get('fields', {}).items()} for d in r.get('documents', [])]
        tok = r.get('nextPageToken')
        if not tok: return docs

def main():
    rows = fetch(token())
    wb = Workbook(); ws = wb.active; ws.title = 'الزوار'; ws.sheet_view.rightToLeft = True
    gold = PatternFill('solid', fgColor='EDBC61'); bold = Font(bold=True, color='1A1208')
    ws.append([a for _, a in FIELDS])
    for c in ws[1]: c.fill = gold; c.font = bold; c.alignment = Alignment(horizontal='center')
    for r in sorted(rows, key=lambda x: x.get('ts', '')): ws.append([r.get(k, '') for k, _ in FIELDS])
    for i, w in enumerate([24,16,20,22,22,34,40,12,14,12,22], 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    s = wb.create_sheet('الملخص'); s.sheet_view.rightToLeft = True
    cnt = lambda key: collections.Counter(r.get(key, '') for r in rows)
    ints = collections.Counter(i for r in rows for i in r.get('interests', '').split('، ') if i)
    lines = [('إجمالي الزوار المسجلين', len(rows)), ('موافقون على المتابعة', sum(r.get('consent') == 'نعم' for r in rows)), ('', '')]
    for title, c in [('حسب اليوم', cnt('day')), ('حسب الفئة', cnt('category')), ('حسب الاهتمام', ints), ('حسب المنظّم', cnt('usher'))]:
        lines += [(title, '')] + [(k or 'غير محدد', v) for k, v in c.most_common()] + [('', '')]
    for a, b in lines:
        s.append([a, b]); print(f'{a:<28} {b}')
    for row in s.iter_rows():
        if row[0].value and row[1].value == '': row[0].fill = gold; row[0].font = bold
    s.column_dimensions['A'].width = 34; s.column_dimensions['B'].width = 10
    out = 'تقرير-زوار-الجناح.xlsx'; wb.save(out); print('->', out, f'({len(rows)} زائر)')

if __name__ == '__main__':
    assert val({'arrayValue': {'values': [{'stringValue': 'أ'}, {'stringValue': 'ب'}]}}) == 'أ، ب'  # self-check
    main()
